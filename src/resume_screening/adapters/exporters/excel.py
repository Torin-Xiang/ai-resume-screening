"""Excel 导出器适配器（方案 §2.1 ⑤、§5）。

目标：产出一张 HR **能直接使用**的表——排名 / 分数 / 命中维度 / 证据引用 / 联系方式 / 原文件路径，
并留两列给 HR 填复核判定（通过 / 淘汰 / 待确认）与备注，供回写。

样式：表头加粗+底色、冻结首行、列宽适配、证据列自动换行；
"待确认"行整行浅黄底（方案 §11 未决问题 4 的处理方式）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from resume_screening.core.models import CandidateResult, Judgement, RunRecord, ScreeningVerdict

VERDICT_TEXT = {
    ScreeningVerdict.PASS: "通过",
    ScreeningVerdict.REJECT: "淘汰",
    ScreeningVerdict.NEEDS_REVIEW: "待确认",
}

HEADERS = [
    "排名",
    "简历文件",
    "姓名",
    "联系方式",
    "初筛判定",
    "得分",
    "满分",
    "命中加分项",
    "未通过硬门槛",
    "待确认项",
    "命中减分项",
    "结论与原文证据",
    "原文件路径",
    "简历ID",
    "HR复核判定",
    "HR备注",
]

WIDTHS = [6, 22, 10, 20, 10, 8, 8, 24, 24, 24, 20, 70, 40, 18, 14, 24]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")  # 待确认：浅黄
REJECT_FILL = PatternFill("solid", fgColor="F2F2F2")  # 淘汰：浅灰
HR_COL_FILL = PatternFill("solid", fgColor="FFF9E6")

HR_DECISION_COL = 15  # O 列
HR_NOTE_COL = 16  # P 列


class ExcelExporter:
    name = "excel"

    def export(self, run: RunRecord, out_path: str | Path) -> Path:
        p = Path(out_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        ws = wb.active
        ws.title = "初筛结果"
        self._write_header(ws)
        for i, r in enumerate(run.results, start=1):
            self._write_row(ws, i + 1, i, r, run)
        ws.freeze_panes = "A2"  # 冻结首行
        self._set_widths(ws)
        self._add_hr_validation(ws, len(run.results))

        self._write_meta_sheet(wb, run)
        wb.save(p)
        return p

    # ---------------- 写入 ----------------
    def _write_header(self, ws) -> None:
        ws.append(HEADERS)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", horizontal="center")

    def _write_row(self, ws, row: int, rank: int, r: CandidateResult, run: RunRecord) -> None:
        ws.cell(row=row, column=1, value=rank)
        ws.cell(row=row, column=2, value=r.file_name)
        ws.cell(row=row, column=3, value=r.candidate_name)
        ws.cell(row=row, column=4, value=r.candidate_contact)
        ws.cell(row=row, column=5, value=VERDICT_TEXT[r.verdict])
        ws.cell(row=row, column=6, value=r.score)
        ws.cell(row=row, column=7, value=run.results[0].max_score if run.results else 0)
        ws.cell(row=row, column=8, value=self._dim_summary(r, "bonus", Judgement.SATISFIED))
        ws.cell(
            row=row,
            column=9,
            value=self._dim_summary(r, "hard_filter", Judgement.NOT_SATISFIED),
        )
        ws.cell(row=row, column=10, value=self._unclear_summary(r))
        ws.cell(row=row, column=11, value=self._dim_summary(r, "red_flag", Judgement.SATISFIED))
        ws.cell(row=row, column=12, value=r.evidence_rows()).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        ws.cell(row=row, column=13, value=r.source_path)
        ws.cell(row=row, column=14, value=r.resume_id)
        ws.cell(row=row, column=HR_DECISION_COL, value=r.hr_decision or "").fill = HR_COL_FILL
        ws.cell(row=row, column=HR_NOTE_COL, value="").fill = HR_COL_FILL

        ws.cell(row=row, column=12).alignment = Alignment(wrap_text=True, vertical="top")
        if r.verdict is ScreeningVerdict.NEEDS_REVIEW:
            for c in range(1, len(HEADERS) + 1):
                if c not in (HR_DECISION_COL, HR_NOTE_COL):
                    ws.cell(row=row, column=c).fill = REVIEW_FILL
        elif r.verdict is ScreeningVerdict.REJECT:
            for c in range(1, len(HEADERS) + 1):
                if c not in (HR_DECISION_COL, HR_NOTE_COL):
                    ws.cell(row=row, column=c).fill = REJECT_FILL

    def _write_meta_sheet(self, wb: Workbook, run: RunRecord) -> None:
        ws = wb.create_sheet("评分卡信息")
        rows = [
            ("run_id", run.run_id),
            ("岗位", run.job_title),
            ("评分卡版本", run.rubric_version),
            ("评分卡路径", run.rubric_path),
            ("简历目录", run.input_dir),
            ("运行时间", run.created_at),
            ("简历份数", len(run.results)),
            ("LLM 调用次数", run.llm_call_count),
            ("LLM token 合计", run.llm_cost_tokens),
            ("", ""),
            ("说明", "每次结果都记录评分卡版本；不同版本的结果不可直接比较（方案 §4.4）"),
            ("回写方式", "在『初筛结果』页填写『HR复核判定』列后执行 review 命令回写"),
        ]
        for k, v in rows:
            ws.append([k, v])
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 70

    @staticmethod
    def _set_widths(ws) -> None:
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    @staticmethod
    def _add_hr_validation(ws, count: int) -> None:
        """HR 复核判定列加下拉：通过/淘汰/待确认（防止手填错别字导致回写失败）。"""
        if count <= 0:
            return
        dv = DataValidation(
            type="list", formula1='"通过,淘汰,待确认"', allow_blank=True, showDropDown=False
        )
        ws.add_data_validation(dv)
        dv.add(f"O2:O{count + 1}")

    # ---------------- 摘要文本 ----------------
    @staticmethod
    def _dim_summary(r: CandidateResult, kind: str, judgement: Judgement) -> str:
        names = [
            f"{d.name}({d.score:+g})" if d.score else d.name
            for d in r.dimensions
            if d.kind == kind and d.judgement is judgement
        ]
        return "、".join(names)

    @staticmethod
    def _unclear_summary(r: CandidateResult) -> str:
        names = [
            d.name
            for d in r.dimensions
            if d.judgement in (Judgement.NOT_MENTIONED, Judgement.UNKNOWN)
        ]
        return "、".join(names)

    # ---------------- 回写读取 ----------------
    def read_decisions(self, path: str | Path) -> dict[str, dict[str, Any]]:
        """从导出的 Excel 读回 HR 判定：key = resume_id。"""
        wb = load_workbook(path, data_only=True)
        ws = wb["初筛结果"] if "初筛结果" in wb.sheetnames else wb.active
        out: dict[str, dict[str, Any]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[13] is None:
                continue
            resume_id = str(row[13])
            decision = (row[HR_DECISION_COL - 1] or "").strip() if len(row) >= HR_DECISION_COL else ""
            note = (row[HR_NOTE_COL - 1] or "").strip() if len(row) >= HR_NOTE_COL else ""
            if decision or note:
                out[resume_id] = {"hr_decision": decision, "hr_note": note}
        return out


def export_run(run: RunRecord, out_path: str | Path) -> Path:
    return ExcelExporter().export(run, out_path)
