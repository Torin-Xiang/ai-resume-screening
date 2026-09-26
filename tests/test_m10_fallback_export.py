"""M10：粘贴文本兜底、导出 Excel（产品方案 §5.1 P0#9、§5.2）。

本文件最核心的一条是 `test_scanned_pdf_gets_a_way_out_via_paste`：
扫描件解析不达标之后，**必须有一条路能让这个人进到排名里**。
产品方案 §6.1 把这条定为硬约束——"绝不允许静默跳过，否则 HR 会误以为
系统看过并放过了这个人"。兜底不是锦上添花，它是"未提及不得淘汰"之外
另一条防止误伤的设计。

其他覆盖：

- 粘贴走的是**同一条链路**（同样的质量阈值、同样调模型），不是简化版
- 同一段文本粘两次不重复计费
- 导出顺序与屏幕一致，且带上 HR 已做的复核判定
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import write_pdf, write_scanned_pdf
from resume_screening.core.models import ParseStatus
from resume_screening.core.quality import assess_text_quality
from resume_screening.store.entities import ResumeStatus
from test_m6_workbench import (
    create_job,
    make_client,
    make_rubric_dict,
    upload,
    wait_settled,
    _resume_items,
)

LONG_TEXT = "\n".join(
    [
        "个人简历",
        "姓名：李示例",
        "联系方式：13800000000",
        "学历：本科",
        "学校：示例大学",
        "专业：计算机科学与技术",
        "工作经历",
        "2023.07-至今 示例科技 算法工程师",
        "负责 RAG 检索增强生成系统的设计与上线，独立完成模型微调。",
    ]
    * 3
)
"""粘贴的简历正文。

注意姓名与联系方式必须和脚本化模型返回的**原文引用**一致
（`EXTRACT_FIELDS` 里 name 的证据是"李示例"）：提取层会校验引用能不能在原文里
找到，对不上就作废为"待确认"。这不是麻烦，正是"引不出原文就承认不知道"的落地。
"""


def paste(client: TestClient, job_id: str, text: str, file_name: str = "李示例"):
    return client.post(
        f"/api/jobs/{job_id}/resumes/text", json={"text": text, "file_name": file_name}
    )


# ——————————————————— 粘贴文本兜底 ———————————————————

def test_scanned_pdf_gets_a_way_out_via_paste(tmp_path: Path) -> None:
    """扫描件 + 粘贴兜底 = 这个人回到排名里。

    这是 P0#9 存在的全部理由。解析不达标的简历如果只能在"需人工"里挂着、
    没有任何出路，HR 就得回到人工翻 PDF 的老路——那产品就白做了。
    """
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        scanned = write_scanned_pdf(tmp_path / "scanned.pdf")
        upload(client, job_id, [scanned])
        wait_settled(client, job_id)

        # 第一份确实进不去排名，且明确说明原因
        before = client.get(f"/api/jobs/{job_id}/ranking").json()
        assert before["total"] == 0
        blocked = client.get(f"/api/jobs/{job_id}/resumes").json()["items"][0]
        assert blocked["status"] == ResumeStatus.NEEDS_MANUAL.value
        assert blocked["error"], "必须说明为什么需要人工处理"

        # HR 粘贴文本后，同一个人立刻参与排名
        assert paste(client, job_id, LONG_TEXT).status_code == 202
        wait_settled(client, job_id)

        after = client.get(f"/api/jobs/{job_id}/ranking").json()

    assert after["total"] == 1
    item = after["groups"][0]["items"][0]
    assert item["candidate_name"] == "李示例"
    assert item["reason"]


def test_pasted_resume_keeps_hr_given_name(tmp_path: Path) -> None:
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        paste(client, job_id, LONG_TEXT, file_name="李示例（Boss直聘）")
        wait_settled(client, job_id)
        items = client.get(f"/api/jobs/{job_id}/resumes").json()["items"]
        ranking = client.get(f"/api/jobs/{job_id}/ranking").json()

    assert items[0]["file_name"] == "李示例（Boss直聘）"
    assert items[0]["origin"] == "paste"
    assert ranking["groups"][0]["items"][0]["file_name"] == "李示例（Boss直聘）"


def test_paste_runs_the_same_pipeline_as_upload(tmp_path: Path) -> None:
    """兜底路径不能是"简化版"：同样要走提取与评分，否则结论与上传的不可比。"""
    with make_client(tmp_path) as (client, llm):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        paste(client, job_id, LONG_TEXT)
        wait_settled(client, job_id)

    assert llm.extract_calls == 1, "粘贴的内容也要做结构化提取"
    assert llm.score_calls == 1, "粘贴的内容也要评分"


def test_paste_applies_the_same_quality_threshold(tmp_path: Path) -> None:
    """同一段内容在两条路径上必须得到同样的质量结论。

    质量自检曾经是解析器里的私有方法；抽到 `core/quality` 就是为了让
    "粘贴文本"复用同一套阈值。否则会出现"PDF 说不行、粘文本说行"，
    而 HR 无从判断到底该信哪个。
    """
    short = "张三 本科"
    quality = assess_text_quality(short)
    assert quality.status is ParseStatus.SUSPECT
    assert quality.reasons, "存疑必须给出原因，不能只是标个状态"

    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        paste(client, job_id, short, file_name="张三")
        wait_settled(client, job_id)
        item = client.get(f"/api/jobs/{job_id}/ranking").json()["groups"][0]["items"][0]

    # 存疑但可用：仍然出结论，只是把"需人工留意"如实带给 HR
    assert item["parse_status"] == ParseStatus.SUSPECT.value


def test_empty_paste_is_rejected(tmp_path: Path) -> None:
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        assert paste(client, job_id, "   ").status_code == 422


def test_same_text_pasted_twice_is_deduplicated(tmp_path: Path) -> None:
    """与上传一致：同一内容不重复计费（产品方案 §2.2）。"""
    with make_client(tmp_path) as (client, llm):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        first = paste(client, job_id, LONG_TEXT).json()
        second = paste(client, job_id, LONG_TEXT).json()
        wait_settled(client, job_id)

        items = client.get(f"/api/jobs/{job_id}/resumes").json()["items"]

    assert first["resume_id"] == second["resume_id"]
    assert second["is_new"] is False
    assert len(items) == 1, "同一段文本不该在列表里出现两次"
    assert llm.extract_calls == 1, "重复粘贴不应重复提取"


def test_paste_without_rubric_is_not_silently_lost(tmp_path: Path) -> None:
    """没有标准时粘贴也要有明确状态，而不是悄悄消失。"""
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client)  # 不设标准
        paste(client, job_id, LONG_TEXT)
        wait_settled(client, job_id)
        items = client.get(f"/api/jobs/{job_id}/resumes").json()["items"]

    assert items[0]["status"] == ResumeStatus.FAILED.value
    assert "标准" in items[0]["error"]


# ——————————————————— 导出 Excel ———————————————————

def test_export_without_rubric_returns_409(tmp_path: Path) -> None:
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client)
        resp = client.get(f"/api/jobs/{job_id}/export")
    assert resp.status_code == 409


def test_export_carries_ranking_and_hr_decisions(tmp_path: Path) -> None:
    """导出的表拿来就能发给业务方：顺序与屏幕一致，且带上 HR 已做的复核。"""
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        pdf = write_pdf(tmp_path / "a.pdf", [*_resume_items()])
        upload(client, job_id, [pdf])
        wait_settled(client, job_id)

        ranking = client.get(f"/api/jobs/{job_id}/ranking").json()
        resume_id = ranking["groups"][0]["items"][0]["resume_id"]
        assert (
            client.post(
                f"/api/jobs/{job_id}/review",
                json={"resume_id": resume_id, "decision": "agree"},
            ).status_code
            == 200
        )

        resp = client.get(f"/api/jobs/{job_id}/export")

    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["初筛结果"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]

    assert row[header.index("简历文件")] == "a.pdf"
    assert row[header.index("姓名")] == "李示例"
    assert row[header.index("HR复核判定")] == "agree", "复核结果必须落到导出表里"
    assert row[header.index("结论与原文证据")], "导出必须带原文证据，否则无法核对"

    meta = wb["评分卡信息"]
    meta_values = {r[0].value: r[1].value for r in meta.iter_rows(min_row=2)}
    assert meta_values.get("评分卡版本") == "v1"


def test_export_order_matches_screen_order(tmp_path: Path) -> None:
    """顺序按分组排（推荐 → 待确认 → 不推荐），与界面一致。

    否则 HR 对着屏幕找第 5 个人，在表里是第 12 行。
    """
    with make_client(tmp_path) as (client, _):
        job_id = create_job(client, rubric=make_rubric_dict("v1"))
        good = write_pdf(tmp_path / "good.pdf", [*_resume_items()])
        # 第二份用同一个脚本化模型也会被判为满足，这里只验证顺序稳定性
        more = write_pdf(tmp_path / "more.pdf", [*_resume_items(offset=400)])
        upload(client, job_id, [good, more])
        wait_settled(client, job_id)

        ranking = client.get(f"/api/jobs/{job_id}/ranking").json()
        expected = [
            item["file_name"]
            for group in ranking["groups"]
            for item in group["items"]
        ]
        resp = client.get(f"/api/jobs/{job_id}/export")

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["初筛结果"]
    header = [c.value for c in ws[1]]
    exported = [row[header.index("简历文件")].value for row in ws.iter_rows(min_row=2)]

    assert exported == expected, "导出顺序必须与界面一致"
    assert len(exported) == 2
