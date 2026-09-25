"""M5：Excel 导出 + 判定回写（方案 §2.1 ⑤、§8 M5）。

验收标准：产出一张能直接使用的表。
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from resume_screening.adapters.exporters.excel import (
    HEADERS,
    HR_DECISION_COL,
    VERDICT_TEXT,
    ExcelExporter,
)
from resume_screening.adapters.llm.fake import FakeLLMClient
from resume_screening.adapters.parsers.pdf_pymupdf import PdfPymupdfParser
from resume_screening.cli.main import app
from resume_screening.core.models import (
    CandidateResult,
    DimensionResult,
    Evidence,
    Judgement,
    ParseStatus,
    RunRecord,
    ScreeningVerdict,
)
from resume_screening.core.pipeline import run_screen
from resume_screening.core.review import apply_decisions
from resume_screening.core.rubric import Bonus, HardFilter, Rubric
from resume_screening.core.run_store import load_run

runner = CliRunner()


def _run() -> RunRecord:
    d1 = DimensionResult(
        dimension_id="bigtech", name="大厂经历", kind="bonus", judgement=Judgement.SATISFIED,
        score=10, reason="在大厂工作过", evidence=[Evidence(quote="示例科技有限公司")], judged_by="llm",
    )
    d2 = DimensionResult(
        dimension_id="ai_exp", name="AI 项目经验", kind="hard_filter",
        judgement=Judgement.NOT_MENTIONED, reason="简历未提及", judged_by="llm",
    )
    return RunRecord(
        run_id="run-test",
        rubric_version="2026-09-26-test",
        rubric_path="rubrics/demo.json",
        job_title="AI 应用开发工程师",
        input_dir="resumes/",
        results=[
            CandidateResult(
                resume_id="id1", file_name="a.pdf", source_path="D:/resumes/a.pdf",
                verdict=ScreeningVerdict.PASS, score=10, max_score=15,
                candidate_name="李示例", candidate_contact="13800000000", dimensions=[d1],
            ),
            CandidateResult(
                resume_id="id2", file_name="b.pdf", source_path="D:/resumes/b.pdf",
                verdict=ScreeningVerdict.NEEDS_REVIEW, score=0, max_score=15,
                candidate_name="王示例", candidate_contact="13900000000", dimensions=[d2],
            ),
            CandidateResult(
                resume_id="id3", file_name="c.pdf", source_path="D:/resumes/c.pdf",
                verdict=ScreeningVerdict.REJECT, score=0, max_score=15, hard_filter_blocked=True,
                candidate_name="赵示例", candidate_contact="13700000000",
                dimensions=[DimensionResult(dimension_id="deg", name="本科及以上", kind="hard_filter",
                                            judgement=Judgement.NOT_SATISFIED, reason="学历为大专",
                                            evidence=[Evidence(quote="学历：大专")], judged_by="rule")],
            ),
        ],
    )


def test_export_creates_usable_sheet(tmp_path: Path) -> None:
    out = ExcelExporter().export(_run(), tmp_path / "out.xlsx")
    wb = load_workbook(out)
    ws = wb["初筛结果"]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == 4  # 表头 + 3 人
    assert ws.freeze_panes == "A2"  # 冻结首行


def test_export_contains_evidence_and_contact(tmp_path: Path) -> None:
    ws = load_workbook(ExcelExporter().export(_run(), tmp_path / "o.xlsx"))["初筛结果"]
    row2 = [c.value for c in ws[2]]
    assert row2[2] == "李示例"
    assert row2[3] == "13800000000"
    assert "示例科技有限公司" in row2[11]  # 结论与原文证据
    assert row2[12] == "D:/resumes/a.pdf"  # 原文件路径
    assert row2[4] == VERDICT_TEXT[ScreeningVerdict.PASS]


def test_export_marks_needs_review_rows(tmp_path: Path) -> None:
    """方案 §11 未决问题 4：待确认在 Excel 中以底色 + 判定列文字呈现。"""
    ws = load_workbook(ExcelExporter().export(_run(), tmp_path / "o.xlsx"))["初筛结果"]
    review_row = [c.value for c in ws[3]]
    assert review_row[4] == "待确认"
    assert ws.cell(row=3, column=1).fill.fgColor.rgb.endswith("FFF2CC")


def test_export_has_hr_dropdown(tmp_path: Path) -> None:
    ws = load_workbook(ExcelExporter().export(_run(), tmp_path / "o.xlsx"))["初筛结果"]
    assert any("通过,淘汰,待确认" in str(dv.formula1) for dv in ws.data_validations.dataValidation)


def test_export_meta_sheet_records_rubric_version(tmp_path: Path) -> None:
    wb = load_workbook(ExcelExporter().export(_run(), tmp_path / "o.xlsx"))
    meta = {row[0]: row[1] for row in wb["评分卡信息"].iter_rows(values_only=True)}
    assert meta["评分卡版本"] == "2026-09-26-test"
    assert meta["run_id"] == "run-test"


def test_rank_ordering(tmp_path: Path) -> None:
    ws = load_workbook(ExcelExporter().export(_run(), tmp_path / "o.xlsx"))["初筛结果"]
    ranks = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    verdicts = [ws.cell(row=r, column=5).value for r in range(2, ws.max_row + 1)]
    assert ranks == [1, 2, 3]
    assert verdicts == ["通过", "待确认", "淘汰"]  # 淘汰排最后


# ---------------- 回写 ----------------
def test_read_and_apply_decisions(tmp_path: Path) -> None:
    out = ExcelExporter().export(_run(), tmp_path / "o.xlsx")
    wb = load_workbook(out)
    ws = wb["初筛结果"]
    ws.cell(row=2, column=HR_DECISION_COL, value="通过")
    ws.cell(row=3, column=HR_DECISION_COL, value="淘汰")
    ws.cell(row=3, column=HR_DECISION_COL + 1, value="经验不足")
    wb.save(out)

    decisions = ExcelExporter().read_decisions(out)
    assert set(decisions) == {"id1", "id2"}
    run = _run()
    updated, unknown, invalid = apply_decisions(run, decisions)
    assert updated == 2 and unknown == [] and invalid == []
    assert run.results[0].hr_decision == "通过"
    assert run.results[1].hr_decision == "淘汰"
    assert any("HR 备注" in n for n in run.results[1].notes)


def test_invalid_and_unknown_decisions_are_reported() -> None:
    run = _run()
    updated, unknown, invalid = apply_decisions(
        run, {"id1": {"hr_decision": "录用"}, "ghost": {"hr_decision": "通过"}}
    )
    assert updated == 0 and unknown == ["ghost"] and invalid == ["id1: 录用"]


# ---------------- 端到端 CLI ----------------
_EXTRACT_PAYLOAD = {
    "fields": {
        "name": {"value": "李示例", "status": "extracted",
                 "evidence": {"quote": "姓名：李示例", "page": 1}},
        "contact": {"value": "13800000000", "status": "extracted",
                    "evidence": {"quote": "联系电话：13800000000", "page": 1}},
        "education": {"value": "本科", "status": "extracted",
                      "evidence": {"quote": "计算机科学与技术 本科", "page": 1}},
        "years_of_experience": {"value": 4, "status": "extracted",
                                "evidence": {"quote": "2022.07-至今", "page": 1}},
    },
    "work_experiences": [{"company": "示例科技", "title": "算法工程师", "start": "2022.07",
                          "end": "至今", "evidence": {"quote": "2022.07-至今", "page": 1}}],
}

_SCORE_PAYLOAD = {
    "results": [
        {"id": "ai_exp", "judgement": "satisfied", "hit_level": 1,
         "reason": "有 RAG 上线经验",
         "evidence": [{"quote": "负责 RAG 检索增强生成系统的设计与上线", "page": 1}]},
        {"id": "bigtech", "judgement": "satisfied", "hit_level": 0.5,
         "reason": "中型公司", "evidence": [{"quote": "示例科技有限公司", "page": 1}]},
    ]
}


def _fake_llm() -> FakeLLMClient:
    return FakeLLMClient(match={"请抽取以下字段": _EXTRACT_PAYLOAD, "请对上面每个维度": _SCORE_PAYLOAD})


def _write_rubric(tmp_path: Path) -> Path:
    r = Rubric(
        version="2026-09-26-e2e",
        job_title="AI 应用开发工程师",
        hard_filters=[
            HardFilter(id="ai_exp", name="AI 项目经验", definition="须实际参与 AI 能力交付",
                       positive_examples=["搭建 RAG 并上线"], negative_examples=["课程 demo"]),
        ],
        bonus=[Bonus(id="bigtech", name="大厂经历", definition="头部公司经历", weight=10,
                     positive_examples=["大厂全职"], negative_examples=["实习"])],
    )
    return r.to_json(tmp_path / "rubric.json")


def test_cli_screen_end_to_end(tmp_path: Path, monkeypatch, single_column_pdf: Path) -> None:
    import resume_screening.cli.screen_cmd as cmd

    monkeypatch.setattr(cmd, "build_llm_client", lambda **kw: _fake_llm())
    resumes = tmp_path / "resumes"
    resumes.mkdir()
    import shutil

    shutil.copy(single_column_pdf, resumes / "a.pdf")

    rubric_path = _write_rubric(tmp_path)
    out = tmp_path / "result.xlsx"
    result = runner.invoke(
        app,
        ["screen", "--rubric", str(rubric_path), "--input", str(resumes),
         "--out", str(out), "--cache-dir", str(tmp_path / "cache"),
         "--run-dir", str(tmp_path / "runs")],
    )
    assert result.exit_code == 0, result.output
    assert out.exists()

    ws = load_workbook(out)["初筛结果"]
    assert ws.max_row == 2
    assert ws.cell(row=2, column=3).value == "李示例"  # 姓名
    assert ws.cell(row=2, column=4).value == "13800000000"  # 联系方式
    assert "RAG" in str(ws.cell(row=2, column=12).value)  # 证据引用


def test_cli_review_roundtrip(tmp_path: Path, monkeypatch, single_column_pdf: Path) -> None:
    import resume_screening.cli.screen_cmd as cmd

    monkeypatch.setattr(cmd, "build_llm_client", lambda **kw: _fake_llm())
    resumes = tmp_path / "resumes"
    resumes.mkdir()
    import shutil

    shutil.copy(single_column_pdf, resumes / "a.pdf")
    rubric_path = _write_rubric(tmp_path)
    out = tmp_path / "result.xlsx"
    runner.invoke(
        app,
        ["screen", "--rubric", str(rubric_path), "--input", str(resumes), "--out", str(out),
         "--cache-dir", str(tmp_path / "cache"), "--run-dir", str(tmp_path / "runs")],
    )

    # HR 填写判定
    wb = load_workbook(out)
    ws = wb["初筛结果"]
    ws.cell(row=2, column=HR_DECISION_COL, value="通过")
    wb.save(out)

    run_json = next((tmp_path / "runs").rglob("run.json"))
    res = runner.invoke(app, ["review", "--excel", str(out), "--run", str(run_json)])
    assert res.exit_code == 0, res.output
    assert "回写完成：1 条" in res.output
    reviewed = load_run(run_json.parent / "review.json")
    assert reviewed.results[0].hr_decision == "通过"


def test_second_screen_reuses_cache(tmp_path: Path, single_column_pdf: Path) -> None:
    """换评分重跑（A/B）时复用解析缓存：不再重复调用模型做提取。"""
    resumes = tmp_path / "resumes"
    resumes.mkdir()
    import shutil

    shutil.copy(single_column_pdf, resumes / "a.pdf")
    rubric_path = _write_rubric(tmp_path)
    rubric = Rubric.from_json(rubric_path)
    calls: list[int] = []

    class CountingLLM:
        name = "counting"

        def __init__(self, inner):
            self._inner = inner

        def complete_json(self, **kw):
            calls.append(1)
            return self._inner.complete_json(**kw)

    llm = CountingLLM(_fake_llm())
    run_screen(
        input_dir=resumes, rubric=rubric, parser=PdfPymupdfParser(), llm=llm,
        exporter=ExcelExporter(), cache_dir=tmp_path / "cache",
        out_path=tmp_path / "r1.xlsx", run_id="run-1",
    )
    first = len(calls)
    run_screen(
        input_dir=resumes, rubric=rubric, parser=PdfPymupdfParser(), llm=llm,
        exporter=ExcelExporter(), cache_dir=tmp_path / "cache",
        out_path=tmp_path / "r2.xlsx", run_id="run-2",
    )
    second = len(calls) - first
    assert second < first, f"第二次应只跑评分不跑提取：first={first}, second={second}"


def test_screen_with_no_resumes_raises(tmp_path: Path) -> None:
    empty = tmp_path / "empty"
    empty.mkdir()
    rubric_path = _write_rubric(tmp_path)
    with __import__("pytest").raises(FileNotFoundError):
        run_screen(
            input_dir=empty, rubric=Rubric.from_json(rubric_path), parser=PdfPymupdfParser(),
            llm=FakeLLMClient(), exporter=ExcelExporter(), cache_dir=tmp_path / "cache",
            out_path=tmp_path / "r.xlsx",
        )
